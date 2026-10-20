#!/usr/bin/env python3
import json
import os
import re
import openpyxl

EXCEL_PATH = "/Users/justin/Desktop/DO/FARFALLINO UPDATED MENU.xlsx"
CLOUDINARY_MAP_PATH = "/Users/justin/Projects/rsc/rsc-web-monorepo/apps/api/src/database/seeds/farfallino-cloudinary-urls.json"
OUTPUT_JSON_PATH = "/Users/justin/Projects/rsc/rsc-web-monorepo/apps/api/src/database/seeds/farfallino-menu.json"

# Typos & Title normalizations
TYPO_REPLACEMENTS = {
    "CAPPUCCION": "Cappuccino",
    "CAFÉ LATTE": "Caffè Latte",
    "FARPPUCCINO": "Farppuccino",
    "TEA BAX SELECTION": "Tea Bag Selection",
    "MALTA GUINESS": "Malta Guinness",
    "ITALAIN WINE": "Italian Wine",
    "FOUNDENTE AL CIOCCOLATO": "Fondente al Cioccolato",
    "FOUNDENTE AL CIOCCOLATO CON GELATO ALLA VANIGLIA": "Fondente al Cioccolato con Gelato alla Vaniglia",
    "CANNOLI ALLA CREMA DI PISTACCHIO": "Cannoli alla Crema di Pistacchio",
    "STRACIATELLA PESTO": "Stracciatella e Pesto",
    "CROCHETTE CON POLLO GAMBERI (spicy)": "Crocchette con Pollo e Gamberi (Spicy)",
    "ALETTE DI POLLO DOLCE PICCANTE (spicy)": "Alette di Pollo Dolce Piccante (Spicy)",
    "SEAMORZA": "Scamorza",
    "ESPRESSO COFFE": "Espresso Coffee",
    "CYNNAMON": "Cinnamon",
    "SCHWEPPES MOJITO": "Schweppes Mojito",
    "JACK DANIEL\"S TENNESSE": "Jack Daniel's Tennessee",
    "MARINI ROSSO": "Martini Rosso",
    "AMARONE DELLA VALPOLICELLA CORTE MANARA 202": "Amarone della Valpolicella Corte Manara 2020",
    "CHIARLI GRASPAROSSA LAMBRUSCO SECO  (DRY)": "Chiarli Grasparossa Lambrusco Secco (Dry)",
    "LAMBRUSCO GRASSPARSA AMABILE SWEET": "Lambrusco Grasparossa Amabile Sweet",
    "FRIULIA VENEZIA GIULIA": "Friuli Venezia Giulia",
    "SETTESOLI BIO GRILLO SICILA  2022": "Settesoli Bio Grillo Sicilia 2022",
    "GRILLO BIO SICILLA DO TENUTE ORESTIADI": "Grillo Bio Sicilia DOC Tenute Orestiadi",
    "PASSIMENTO BIANCO BAGLIO GIBELLINA  2022": "Passimento Bianco Baglio Gibellina 2022",
    "CANTINE FINA MIRAL GRILLO": "Cantine Fina Miral Grillo",
    "TENNESSEE HONEY SMASH": "Tennessee Honey Smash",
    "MEDITERRANEAN NEGRONI": "Mediterranean Negroni",
    "LA DOLCE VITA": "La Dolce Vita",
    "LUNA DI MIELE": "Luna di Miele",
    "COSTA AZZURRA": "Costa Azzurra",
    "E TONICO GIN": "È Tonico Gin",
    "PROSECCO SUPERIORE VALDOBBIADENE0": "Prosecco Superiore Valdobbiadene",
}

DESC_TYPO_REPLACEMENTS = {
    "SEAMORZA": "scamorza",
    "COFFE": "coffee",
    "CYNNAMON": "cinnamon",
    "TENNESSE": "Tennessee",
    "JACK DANIEL\"S": "Jack Daniel's",
    "ARRABIATTA": "arrabbiata",
    "ROCCA SALAD": "rocket salad",
}

def clean_title(raw: str) -> str:
    if not raw:
        return ""
    s = raw.strip()
    for typo, fix in TYPO_REPLACEMENTS.items():
        if typo.lower() == s.lower():
            return fix
    
    # Title Case conversion
    words = s.split()
    capitalized = []
    lower_connectors = {
        "di", "del", "della", "delle", "dei", "degli", "da", "dal", "dalla",
        "con", "e", "ed", "al", "alla", "alle", "ai", "agli", "in", "per", "fra", "tra",
        "o", "od", "su", "sul", "sulla", "the", "a", "an", "and", "or", "of", "with"
    }
    for idx, w in enumerate(words):
        lw = w.lower()
        if idx > 0 and lw in lower_connectors and not (w.startswith("(") or w.endswith(")")):
            capitalized.append(lw)
        else:
            # Keep letters capitalized after apostrophe (e.g. d'Alba -> d'Alba, Dell'Aglio)
            if "'" in w:
                parts = w.split("'")
                fixed_parts = [p.capitalize() for p in parts]
                if fixed_parts[0].lower() in ["d", "l", "un", "all"]:
                    fixed_parts[0] = fixed_parts[0].lower()
                capitalized.append("'".join(fixed_parts))
            else:
                capitalized.append(w.capitalize())
    res = " ".join(capitalized)
    # Fix specific abbreviations
    res = re.sub(r'\bYo\b', 'YO', res, flags=re.I)
    res = re.sub(r'\bDoc\b', 'DOC', res, flags=re.I)
    res = re.sub(r'\bDocg\b', 'DOCG', res, flags=re.I)
    res = re.sub(r'\bXo\b', 'XO', res, flags=re.I)
    res = re.sub(r'\bVsop\b', 'VSOP', res, flags=re.I)
    res = re.sub(r'\bMl\b', 'ml', res, flags=re.I)
    res = re.sub(r'\bLtr\b', 'L', res, flags=re.I)
    res = re.sub(r'\b1\.5ltr\b', '1.5L', res, flags=re.I)
    return res

def clean_desc(raw: str) -> str:
    if not raw:
        return ""
    s = raw.strip()
    for typo, fix in DESC_TYPO_REPLACEMENTS.items():
        s = re.sub(re.escape(typo), fix, s, flags=re.IGNORECASE)
    
    # Ensure nice punctuation spacing: replace comma with comma-space if no space follows
    s = re.sub(r',([^\s])', r', \1', s)
    s = re.sub(r'\s+', ' ', s)

    # Sentence casing
    if len(s) > 1:
        s = s[0].upper() + s[1:].lower()
        # Capitalize after periods
        s = re.sub(r'(\.\s+[a-z])', lambda m: m.group(1).upper(), s)
        # Capitalize known proper nouns in description
        proper_nouns = [
            "Italian", "Macallan", "Jack Daniel's", "Tennessee", "Gin Mare", "Campari", "Schweppes",
            "Bacardi", "Ciroc", "Grey Goose", "Belvedere", "Hennessy", "Martell", "Remy Martin",
            "Martini", "Moët", "Veuve Clicquot", "Dom Pérignon", "Tignanello", "Solaia", "Chianti",
            "Prosecco", "Spumante", "Bufala", "Mozzarella", "Parmigiano", "Stracciatella", "Arrabbiata",
            "Bolognese", "Alfredo", "Carbonara", "Puttanesca", "Genovese", "Romana", "Toscana",
            "Pugliese", "Calabrese"
        ]
        for pn in proper_nouns:
            s = re.sub(r'\b' + re.escape(pn) + r'\b', pn, s, flags=re.IGNORECASE)
    return s.strip()

def parse_price_minor(row_num: int, raw_val) -> int:
    if raw_val is None:
        return 0
    if isinstance(raw_val, str):
        val = float(raw_val.replace(',', '.'))
    else:
        val = float(raw_val)

    if row_num < 163:
        # Food & cafe & drinks section 1
        naira = val * 1000
    elif row_num >= 301:
        # Wines section
        if val < 5:
            naira = val * 1_000_000
        else:
            naira = val * 1000
    else:
        # Explicit drinks / champagne / spirits section
        naira = val

    return int(round(naira * 100))

def main():
    cloudinary_urls = {}
    if os.path.exists(CLOUDINARY_MAP_PATH):
        try:
            with open(CLOUDINARY_MAP_PATH, "r") as f:
                cloudinary_urls = json.load(f)
        except Exception as e:
            print(f"Error loading Cloudinary map: {e}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    categories_list = []
    current_cat = None
    cat_items = {}
    current_region = None
    is_french_section = False

    for r in range(4, ws.max_row + 1):
        c2 = ws.cell(row=r, column=2).value
        c6 = ws.cell(row=r, column=6).value
        c9 = ws.cell(row=r, column=9).value
        c16 = ws.cell(row=r, column=16).value
        c19 = ws.cell(row=r, column=19).value

        c2_str = str(c2).strip() if c2 else ''
        c6_str = str(c6).strip() if c6 else ''
        c9_str = str(c9).strip() if c9 else ''

        if 'FRENCH WINE' in c2_str:
            is_french_section = True

        # Header detection
        if c2_str and c2_str not in [
            'CATEGORY', 'FARFALLINO MENU', 'DRINK MENU',
            'ALL HAM CAN BE REPLACED BY SMOKED TURKEY', 'OR BEEF PEPPERONI',
            'CHOOSE ONE SIDE : ROASTED POTATO,', 'FRENCH FRIES,VEGETABLE STIRFRY RICE,',
            'LINGUINE SPICY TOMATO OR SAUTEED VEGETABLES'
        ]:
            if 'ANTIPASTI' in c2_str: current_cat = 'Antipasti & Starters'
            elif 'CRÈME' in c2_str or 'SOUP' in c2_str: current_cat = 'Zuppe & Minestre (Soups)'
            elif 'FOCACCIA' in c2_str: current_cat = 'Focaccia'
            elif 'INSALATE' in c2_str: current_cat = 'Insalate (Salads)'
            elif 'BURRATA' in c2_str: current_cat = 'Burrata'
            elif 'PRIMI' in c2_str or 'PASTA' in c2_str: current_cat = 'Primi Piatti (Pasta)'
            elif 'RISOTTI' in c2_str: current_cat = 'Risotti'
            elif 'SECONDI' in c2_str: current_cat = 'Secondi Piatti (Meat & Fish)'
            elif 'CONTORNI' in c2_str: current_cat = 'Contorni (Side Dishes)'
            elif 'PANUOZZO' in c2_str: current_cat = 'Panuozzo & Puccia'
            elif 'PIZZE' in c2_str or 'PIZZA' in c2_str: current_cat = 'Pizze (Pizzas)'
            elif 'DOLCI' in c2_str: current_cat = 'Dolci (Desserts)'
            elif 'MILKSHAKE' in c2_str: current_cat = 'Milkshakes'
            elif 'CAFFE' in c2_str: current_cat = 'Caffè & Hot Drinks'
            elif 'BEVANDE' in c2_str: current_cat = 'Bevande (Beverages)'
            elif 'SUCCO' in c2_str: current_cat = 'Succhi (Juices)'
            elif 'BIRRA' in c2_str: current_cat = 'Birre (Beers)'
            elif 'FARFALLINO SIGNATURE' in c2_str: current_cat = 'Signature Cocktails'
            elif 'SPRITZ' in c2_str: current_cat = 'Spritz'
            elif 'COCKTAIL' in c2_str: current_cat = 'Cocktails'
            elif 'MOCKTAIL' in c2_str: current_cat = 'Mocktails'
            elif 'GRAPPA' in c2_str: current_cat = 'Grappa'
            elif 'DIGESTIVE' in c2_str: current_cat = 'Digestivi (Digestifs & Liqueurs)'
            elif 'WHISKY' in c2_str: current_cat = 'Whisky & Whiskey'
            elif 'TEQUILA' in c2_str: current_cat = 'Tequila & Mezcal'
            elif 'GIN' in c2_str: current_cat = 'Gin'
            elif 'VODKA' in c2_str: current_cat = 'Vodka'
            elif 'COGNAC' in c2_str: current_cat = 'Cognac'
            elif 'VERMUT' in c2_str or 'VERMOUTH' in c2_str: current_cat = 'Vermut & Aperitivi'
            elif 'CHAMPAGNE' in c2_str: current_cat = 'Champagne'
            elif 'WINE PROSECCO & SPUMANTE' in c2_str: current_cat = 'Vini al Calice (Wine by the Glass)'
            elif 'WHITE WINE' in c2_str:
                current_cat = 'Vini Bianchi Francesi (French White Wines)' if is_french_section else 'Vini Bianchi Italiani (Italian White Wines)'
            elif 'ROSE WINE' in c2_str:
                current_cat = 'Vini Rosati Francesi (French Rosé Wines)' if is_french_section else 'Vini Rosati Italiani (Italian Rosé Wines)'
            elif 'RED WINE' in c2_str:
                current_cat = 'Vini Rossi Francesi (French Red Wines)' if is_french_section else 'Vini Rossi Italiani (Italian Red Wines)'
            elif 'PROSECCO' in c2_str: current_cat = 'Prosecco'
            elif 'SPUMANTE' in c2_str: current_cat = 'Spumante'

            if current_cat and current_cat not in categories_list:
                categories_list.append(current_cat)

        # Skip rows that are modifier rows
        if r in [16, 113, 114, 115, 116]:
            continue

        title = None
        desc = None

        if current_cat in ['Vini Bianchi Italiani (Italian White Wines)', 'Vini Rosati Italiani (Italian Rosé Wines)', 'Vini Rossi Italiani (Italian Red Wines)']:
            if c6_str: current_region = c6_str
            title = c9_str
            desc = f"Region: {clean_title(current_region)}" if current_region else ""
        elif current_cat == 'Vini al Calice (Wine by the Glass)':
            title = c9_str or c6_str
            desc = "By the glass"
        elif current_cat in ['Whisky & Whiskey', 'Tequila & Mezcal', 'Gin', 'Vodka', 'Cognac'] and c9 and isinstance(c9, (int, float)):
            # Shot price in c9, bottle price in c16
            title = c6_str
            shot_price = int(c9)
            bottle_price = int(c16) if c16 else 0
            desc = f"Bottle: ₦{bottle_price:,} | Shot / Glass: ₦{shot_price:,}"
        else:
            title = c6_str or c9_str
            desc = c9_str if c6_str and c9_str != c6_str else ""

        if not title or c16 is None:
            continue

        item_title = clean_title(title)
        item_desc = clean_desc(desc)
        price_minor = parse_price_minor(r, c16)

        # Image lookup
        image_url = None
        if c19:
            raw_img = str(c19).strip()
            fname = os.path.basename(raw_img).lower()
            if fname in cloudinary_urls:
                image_url = cloudinary_urls[fname]
            elif raw_img.lower() in cloudinary_urls:
                image_url = cloudinary_urls[raw_img.lower()]

        # Modifier groups
        modifier_group_refs = []
        if current_cat == 'Secondi Piatti (Meat & Fish)':
            modifier_group_refs.append('Secondi Side Choice')
        elif current_cat == 'Pizze (Pizzas)':
            modifier_group_refs.append('Pizza Meat Replacement')
            modifier_group_refs.append('Pizza Extra Toppings')
        elif 'cheesecake' in item_title.lower() and 'italiana' in item_title.lower():
            modifier_group_refs.append('Cheesecake Sauce Choice')

        cat_items.setdefault(current_cat, []).append({
            "name": item_title,
            "description": item_desc or None,
            "priceMinor": price_minor,
            "imageUrl": image_url,
            "isAvailable": True,
            "modifierGroupRefs": modifier_group_refs
        })

    # Build modifier groups
    modifier_groups = [
        {
            "name": "Secondi Side Choice",
            "minSelections": 1,
            "maxSelections": 1,
            "isRequired": True,
            "modifiers": [
                { "name": "Roasted Potato", "priceDeltaMinor": 0 },
                { "name": "French Fries", "priceDeltaMinor": 0 },
                { "name": "Vegetable Stirfry Rice", "priceDeltaMinor": 0 },
                { "name": "Linguine Spicy Tomato", "priceDeltaMinor": 0 },
                { "name": "Sautéed Vegetables", "priceDeltaMinor": 0 }
            ]
        },
        {
            "name": "Pizza Meat Replacement",
            "minSelections": 0,
            "maxSelections": 1,
            "isRequired": False,
            "modifiers": [
                { "name": "Replace Ham with Smoked Turkey", "priceDeltaMinor": 0 },
                { "name": "Replace Ham with Beef Pepperoni", "priceDeltaMinor": 0 }
            ]
        },
        {
            "name": "Pizza Extra Toppings",
            "minSelections": 0,
            "maxSelections": 5,
            "isRequired": False,
            "modifiers": [
                { "name": "Shrimps", "priceDeltaMinor": 800000 },
                { "name": "Fresh Mozzarella", "priceDeltaMinor": 600000 },
                { "name": "Turkey", "priceDeltaMinor": 600000 },
                { "name": "Beef Pepperoni", "priceDeltaMinor": 600000 }
            ]
        },
        {
            "name": "Cheesecake Sauce Choice",
            "minSelections": 1,
            "maxSelections": 1,
            "isRequired": True,
            "modifiers": [
                { "name": "Caramel Sauce", "priceDeltaMinor": 0 },
                { "name": "Chocolate Sauce", "priceDeltaMinor": 0 },
                { "name": "Strawberry Sauce", "priceDeltaMinor": 0 }
            ]
        }
    ]

    # Package seed data
    output_data = {
        "outletId": "bcbe3805-9bb5-42d6-b6aa-248010d64385",
        "outletName": "Farfallino",
        "modifierGroups": modifier_groups,
        "categories": [
            {
                "name": cat_name,
                "sortOrder": idx,
                "isActive": True,
                "items": [
                    { **item, "sortOrder": i }
                    for i, item in enumerate(cat_items.get(cat_name, []))
                ]
            }
            for idx, cat_name in enumerate(categories_list)
            if cat_name in cat_items and len(cat_items[cat_name]) > 0
        ]
    }

    os.makedirs(os.path.dirname(OUTPUT_JSON_PATH), exist_ok=True)
    with open(OUTPUT_JSON_PATH, "w") as f:
        json.dump(output_data, f, indent=2)

    total_items = sum(len(c["items"]) for c in output_data["categories"])
    items_with_images = sum(
        1 for c in output_data["categories"] for it in c["items"] if it["imageUrl"] is not None
    )

    print(f"✅ Generated Farfallino seed data:")
    print(f"   Categories: {len(output_data['categories'])}")
    print(f"   Total items: {total_items}")
    print(f"   Items with images: {items_with_images}")
    print(f"   Modifier groups: {len(modifier_groups)}")
    
    # Also write TypeScript module
    ts_path = OUTPUT_JSON_PATH.replace('.json', '.data.ts')
    with open(ts_path, 'w') as f:
        f.write('// Auto-generated Farfallino menu seed data\n')
        f.write('export interface FarfallinoSeedData {\n')
        f.write('  outletId: string;\n')
        f.write('  outletName: string;\n')
        f.write('  modifierGroups: Array<{\n')
        f.write('    name: string;\n')
        f.write('    minSelections: number;\n')
        f.write('    maxSelections: number;\n')
        f.write('    isRequired: boolean;\n')
        f.write('    modifiers: Array<{ name: string; priceDeltaMinor: number }>;\n')
        f.write('  }>;\n')
        f.write('  categories: Array<{\n')
        f.write('    name: string;\n')
        f.write('    sortOrder: number;\n')
        f.write('    isActive: boolean;\n')
        f.write('    items: Array<{\n')
        f.write('      name: string;\n')
        f.write('      description: string | null;\n')
        f.write('      priceMinor: number;\n')
        f.write('      imageUrl: string | null;\n')
        f.write('      isAvailable: boolean;\n')
        f.write('      sortOrder: number;\n')
        f.write('      modifierGroupRefs: string[];\n')
        f.write('    }>;\n')
        f.write('  }>;\n')
        f.write('}\n\n')
        f.write('export const farfallinoMenuData: FarfallinoSeedData = ')
        json.dump(output_data, f, indent=2)
        f.write(';\n')

    print(f"   Written to: {OUTPUT_JSON_PATH} and {ts_path}")

if __name__ == "__main__":
    main()
